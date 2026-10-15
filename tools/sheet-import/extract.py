"""Extract every ledger entry from each spreadsheet snapshot into a JSON cache.

Usage: python extract.py SNAPSHOT_DIR OUT_JSON [EXTRA_SNAPSHOT.xlsx ...]

Layout of the ledger tabs (Cantina_Receber):
  * The sheet is a set of vertical *lanes*, each three columns wide:
        date | item | amount
    (lanes are sometimes separated by an empty column).
  * A lane holds several customers stacked vertically. A row with text but no
    amount and no date is a customer name and starts a new block; row 1 is
    just the first such name. The name may sit in the date or the item column.
    Notes inside a block (".", "REFRI?", "SEM JANTA", "1kcafé") don't start one.
  * Rows without a date inherit the last date above them in the same block.
  * Amount cells may be formulas: `=2*8.5` is an item, `=SUM(...)` is a
    running subtotal and is skipped.
"""
import datetime
import json
import os
import re
import sys
from collections import Counter

import openpyxl

LEDGER_TABS = {'PENDENTES', 'RESIDENTES', 'ESTAGIARIOS', '2024', '9O.SEM', 'INOVA',
               'BOLSISTAS', '2025', '2026', 'CALOTE'}
MONTHS = {m: i for i, m in enumerate(
    ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto',
     'setembro', 'outubro', 'novembro', 'dezembro'], 1)}
FILE_RX = re.compile(r'(\d{1,2}) de (\w+)(?: de (\d{4}))?, (\d{1,2})_(\d{2})')
CURRENT_RX = re.compile(r'current_(\d{4}-\d{2}-\d{2})_(\d{2})-(\d{2})')
DATE_TEXT_RX = re.compile(r'(?<!\d)(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{2,4}))?(?!\d)')
MONEY_TEXT_RX = re.compile(r'^\s*-?\s*(r\$)?\s*-?\d+([.,]\d{1,2})?\s*$', re.I)
SUBTOTAL_RX = re.compile(r'sum\(|:', re.I)
EXCEL_EPOCH = datetime.date(1899, 12, 30)
# Column headings some blocks have under the customer name ("data | produto | valor").
HEADING_WORDS = {'data', 'dia', 'produto', 'produtos', 'item', 'itens', 'descricao', 'descrição',
                 'valor', 'vlr', 'preco', 'preço', 'total', 'quantidade', 'qtd'}


def snapshot_time(filename):
    m = FILE_RX.search(filename)
    if m:
        d, mon, y, hh, mm = m.groups()
        return datetime.datetime(int(y or 2026), MONTHS[mon.lower()], int(d), int(hh), int(mm))
    m = CURRENT_RX.search(filename)
    if m:
        return datetime.datetime.fromisoformat(f'{m.group(1)}T{m.group(2)}:{m.group(3)}')
    if 'mais antigo' in filename:
        return datetime.datetime(2024, 1, 1)  # oldest export; sorts before every dated copy
    raise ValueError(f'cannot read a date from file name: {filename}')


def as_date(v, snap_date):
    """Date-column cell -> date or None. Accepts real dates, serial numbers and dd/mm[/yy] text."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 40000 <= v <= 50000:
        return EXCEL_EPOCH + datetime.timedelta(days=int(v))
    if isinstance(v, str) and len(v.strip()) <= 14:  # short text only, tolerates typos like "F24/11"
        m = DATE_TEXT_RX.search(v)
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
            year = (2000 + int(y) if len(y) == 2 else int(y)) if y else snap_date.year
            try:
                out = datetime.date(year, mo, d)
            except ValueError:
                return None
            if not y and out > snap_date:
                out = out.replace(year=year - 1)
            return out
    return None


def as_amount(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return round(float(v), 2)
    if isinstance(v, str) and MONEY_TEXT_RX.match(v):
        s = v.lower().replace('r$', '').replace(' ', '').replace(',', '.')
        try:
            return round(float(s), 2)
        except ValueError:
            return None
    return None


def text(v):
    return v.strip() if isinstance(v, str) and v.strip() and not v.startswith('=') else ''


def cell(grid, r, c):
    if r < 0 or c < 0 or r >= len(grid) or c >= len(grid[r]):
        return None
    return grid[r][c]


def is_entry(values, r, c):
    """An entry is an amount with an item label directly to its left."""
    return as_amount(cell(values, r, c)) is not None and bool(text(cell(values, r, c - 1)))


def is_note(label, date_text):
    """Text rows inside a block that are notes, not a new customer: ".", "REFRI?", "SEM JANTA", "1kcafé"."""
    t = (label or date_text).lower()
    return (not re.search(r'[a-zà-ú]', t) or t.endswith('?') or t.startswith('sem ')
            or bool(label and label[0].isdigit()))


def owners_by_row(values, item_col, snap_date, known_items):
    """Scan one lane top to bottom. For each row, the customer name above it
    and the nearest date since that name: {row: (name, name_row, date)}."""
    out, name, name_row, last_date = {}, None, None, None
    for r in range(len(values)):
        d = as_date(cell(values, r, item_col - 1), snap_date)
        out[r] = (name, name_row, d or last_date)
        has_amount = as_amount(cell(values, r, item_col + 1)) is not None
        label = text(cell(values, r, item_col))
        date_text = text(cell(values, r, item_col - 1)) if not d else ''
        heading = {label.lower(), date_text.lower(), text(cell(values, r, item_col + 1)).lower()} & HEADING_WORDS
        if heading:
            continue
        if not has_amount and not d and (label or date_text):
            if is_note(label, date_text) and name:
                continue
            if not (label and label.lower() in known_items and not date_text):
                name, name_row, last_date = label or date_text, r + 1, None
                continue
        if d:
            last_date = d
    return out


def extract_sheet(values, formulas, snap_date, known_items):
    entries, skipped, lanes = [], Counter(), {}
    for r, row in enumerate(values):
        for c in range(1, len(row)):
            if not is_entry(values, r, c):
                continue
            formula = cell(formulas, r, c)
            if isinstance(formula, str) and formula.startswith('=') and SUBTOTAL_RX.search(formula):
                skipped['subtotal formula'] += 1
                continue
            if c - 1 not in lanes:
                lanes[c - 1] = owners_by_row(values, c - 1, snap_date, known_items)
            customer, name_row, date = lanes[c - 1][r]
            entries.append({
                'customer': customer, 'block_row': name_row, 'row': r + 1,
                'date': date.isoformat() if date else None,
                'label': text(cell(values, r, c - 1)), 'amount': as_amount(cell(values, r, c)),
                'lane': openpyxl.utils.get_column_letter(c),  # item column
            })
    return entries, skipped


def load_grid(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def extract_file(path):
    snap = snapshot_time(os.path.basename(path))
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    wf = openpyxl.load_workbook(path, read_only=True)
    sheets = [(ws.title, load_grid(ws), load_grid(wf[ws.title]))
              for ws in wv.worksheets if ws.title.strip().upper() in LEDGER_TABS]
    wv.close()
    wf.close()

    # Item labels seen with an amount: a lone "salgado" row is an item missing its price, not a customer.
    counts = Counter()
    for _, values, _ in sheets:
        for r, row in enumerate(values):
            for c in range(1, len(row)):
                if is_entry(values, r, c):
                    counts[text(row[c - 1]).lower()] += 1
    known = {k for k, n in counts.items() if n >= 5}

    entries, skipped = [], Counter()
    for title, values, formulas in sheets:
        es, sk = extract_sheet(values, formulas, snap.date(), known)
        for e in es:
            e['tab'] = title.strip()
        entries += es
        skipped.update(sk)
    return {'file': os.path.basename(path), 'time': snap.isoformat(), 'entries': entries, 'skipped': dict(skipped)}


def main(folder, out_path, *extra):
    paths = [os.path.join(folder, n) for n in os.listdir(folder) if n.endswith('.xlsx')] + list(extra)
    paths.sort(key=lambda p: snapshot_time(os.path.basename(p)))
    snaps = []
    for p in paths:
        s = extract_file(p)
        snaps.append(s)
        print(f"{s['time'][:16]} {len(s['entries']):6} entries  skipped={s['skipped']}  {s['file'][-40:]}", flush=True)
    with open(out_path, 'w') as f:
        json.dump(snaps, f)


if __name__ == '__main__':
    main(*sys.argv[1:])
